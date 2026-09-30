"""Master resume (profile/resume.yaml): load/save, text extraction, and .docx rendering."""
from __future__ import annotations

import re
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pydantic import BaseModel

from .config import MASTER_RESUME, OUTPUT_DIR


class SkillGroup(BaseModel):
    category: str
    items: list[str]


class Experience(BaseModel):
    company: str
    title: str
    location: str
    start: str
    end: str
    bullets: list[str]


class Education(BaseModel):
    degree: str
    school: str
    year: str


class Resume(BaseModel):
    name: str
    headline: str
    email: str
    phone: str
    location: str
    links: list[str]
    summary: str
    skills: list[SkillGroup]
    experience: list[Experience]
    education: list[Education]
    certifications: list[str]


def load_master(path: Path = MASTER_RESUME) -> Resume:
    if not path.exists():
        raise SystemExit(
            f"No master resume at {path}.\n"
            "Run: python -m jobbot import-resume <your resume .pdf/.docx>"
        )
    with open(path, encoding="utf-8") as f:
        return Resume.model_validate(yaml.safe_load(f))


def save_master(resume: Resume, path: Path = MASTER_RESUME) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        yaml.safe_dump(resume.model_dump(), f, sort_keys=False, allow_unicode=True, width=100)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return path.read_text(encoding="utf-8", errors="ignore")


def resume_text(r: Resume) -> str:
    """Flat text of the whole resume, used for the anti-fabrication checks."""
    parts = [r.headline, r.summary]
    parts += [f"{g.category}: {', '.join(g.items)}" for g in r.skills]
    for e in r.experience:
        parts += [e.company, e.title, *e.bullets]
    parts += [f"{e.degree} {e.school}" for e in r.education] + r.certifications
    return "\n".join(parts)


def _slug(text: str, n: int = 40) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")[:n] or "job"


def render_docx(r: Resume, company: str, title: str, job_key: str) -> Path:
    """Write a clean one-column resume and return its path."""
    out_dir = OUTPUT_DIR / "resumes"
    out_dir.mkdir(parents=True, exist_ok=True)
    first_last = _slug(r.name, 30)
    path = out_dir / f"{first_last}_{_slug(company, 25)}_{_slug(title, 35)}_{job_key}.docx"

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(40)
        s.left_margin = s.right_margin = Pt(50)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(2)

    def heading(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(r.name)
    run.bold = True
    run.font.size = Pt(18)
    if r.headline:
        p = doc.add_paragraph(r.headline)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact = " | ".join(x for x in [r.location, r.phone, r.email, *r.links] if x)
    p = doc.add_paragraph(contact)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if r.summary:
        heading("Professional Summary")
        doc.add_paragraph(r.summary)

    if r.skills:
        heading("Technical Skills")
        for g in r.skills:
            p = doc.add_paragraph()
            p.add_run(f"{g.category}: ").bold = True
            p.add_run(", ".join(g.items))

    if r.experience:
        heading("Professional Experience")
        for e in r.experience:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.add_run(f"{e.title}").bold = True
            p.add_run(f" | {e.company}" + (f", {e.location}" if e.location else ""))
            p.add_run(f"\t{e.start} - {e.end}").italic = True
            for b in e.bullets:
                doc.add_paragraph(b, style="List Bullet")

    if r.education:
        heading("Education")
        for ed in r.education:
            doc.add_paragraph(", ".join(x for x in [ed.degree, ed.school, ed.year] if x))

    if r.certifications:
        heading("Certifications")
        for c in r.certifications:
            doc.add_paragraph(c, style="List Bullet")

    doc.save(str(path))
    return path
